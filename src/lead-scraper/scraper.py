"""
Lead Scraper - Find local businesses on Google Maps and save to Excel.
Usage: python scraper.py "plumbers" "Hammond Louisiana"
"""

import sys
import re
import time
from pathlib import Path

try:
    from playwright.sync_api import sync_playwright
except ImportError:
    print("Installing playwright...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "playwright"])
    subprocess.check_call([sys.executable, "-m", "playwright", "install", "chromium"])
    from playwright.sync_api import sync_playwright

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def scrape_google_maps(business_type: str, city: str) -> list[dict]:
    """Scrape Google Maps for businesses and return a list of results."""
    query = f"{business_type} in {city}"
    url = f"https://www.google.com/maps/search/{query.replace(' ', '+')}"

    print(f"\nSearching Google Maps for: {query}")
    print(f"URL: {url}\n")

    results = []

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1280, "height": 900},
        )
        page = context.new_page()

        try:
            page.goto(url, wait_until="networkidle", timeout=30000)
        except Exception:
            # networkidle can be flaky, just wait a bit
            page.goto(url, timeout=30000)

        time.sleep(3)

        # Accept cookies/consent if prompted
        try:
            consent = page.locator("button:has-text('Accept all')")
            if consent.count() > 0:
                consent.first.click()
                time.sleep(1)
        except Exception:
            pass

        # Scroll the results panel to load more listings
        feed = page.locator('div[role="feed"]')
        if feed.count() == 0:
            # Try alternative selector
            feed = page.locator('div[role="main"]')

        print("Loading results", end="", flush=True)
        for _ in range(8):
            try:
                feed.evaluate("el => el.scrollTop = el.scrollHeight")
            except Exception:
                try:
                    page.mouse.wheel(0, 3000)
                except Exception:
                    pass
            time.sleep(1.5)
            print(".", end="", flush=True)
        print(" done!\n")

        # Find all business listing links
        links = page.locator('a[href*="/maps/place/"]')
        count = links.count()
        print(f"Found {count} listings. Extracting details...\n")

        for i in range(count):
            try:
                link = links.nth(i)
                # Get the aria-label which usually has the business name
                label = link.get_attribute("aria-label") or ""
                if not label:
                    continue

                # Click into the listing to get details
                link.click()
                time.sleep(2)

                name = label.strip()
                phone = ""
                address = ""
                has_website = "No"

                # Extract info from the detail panel
                # Phone number
                try:
                    phone_el = page.locator(
                        'button[data-item-id*="phone:"] div.fontBodyMedium, '
                        'button[aria-label*="Phone:"]'
                    )
                    if phone_el.count() > 0:
                        phone_text = phone_el.first.text_content() or ""
                        phone_label = (
                            page.locator('button[data-item-id*="phone:"]')
                            .first.get_attribute("aria-label") or ""
                        )
                        # Try aria-label first (more reliable)
                        if phone_label:
                            match = re.search(r"[\d\(\)\-\+\s\.]{7,}", phone_label)
                            if match:
                                phone = match.group().strip()
                        if not phone and phone_text:
                            match = re.search(r"[\d\(\)\-\+\s\.]{7,}", phone_text)
                            if match:
                                phone = match.group().strip()
                except Exception:
                    pass

                # Address
                try:
                    addr_el = page.locator(
                        'button[data-item-id="address"] div.fontBodyMedium, '
                        'button[aria-label*="Address:"]'
                    )
                    if addr_el.count() > 0:
                        addr_label = (
                            page.locator('button[data-item-id="address"]')
                            .first.get_attribute("aria-label") or ""
                        )
                        if addr_label:
                            address = addr_label.replace("Address: ", "").strip()
                        else:
                            address = addr_el.first.text_content().strip()
                except Exception:
                    pass

                # Website
                try:
                    website_el = page.locator(
                        'a[data-item-id="authority"], '
                        'button[data-item-id="authority"]'
                    )
                    if website_el.count() > 0:
                        has_website = "Yes"
                except Exception:
                    pass

                result = {
                    "name": name,
                    "phone": phone,
                    "address": address,
                    "website": has_website,
                }
                results.append(result)

                status = "HAS WEBSITE" if has_website == "Yes" else "no website"
                print(f"  [{i+1}/{count}] {name} — {status}")

                # Go back to results list
                try:
                    back_btn = page.locator('button[aria-label="Back"]')
                    if back_btn.count() > 0:
                        back_btn.first.click()
                        time.sleep(1.5)
                except Exception:
                    page.go_back()
                    time.sleep(2)

            except Exception as e:
                print(f"  [{i+1}/{count}] Skipped (error: {e})")
                try:
                    page.go_back()
                    time.sleep(1.5)
                except Exception:
                    pass
                continue

        browser.close()

    return results


def save_to_excel(results: list[dict], business_type: str, city: str) -> str:
    """Save results to a styled Excel file and return the file path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Styles
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    yes_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    no_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # Headers
    headers = ["Business Name", "Phone Number", "Address", "Website"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Data rows
    for row_idx, biz in enumerate(results, 2):
        ws.cell(row=row_idx, column=1, value=biz["name"]).border = thin_border
        ws.cell(row=row_idx, column=2, value=biz["phone"]).border = thin_border
        ws.cell(row=row_idx, column=3, value=biz["address"]).border = thin_border
        website_cell = ws.cell(row=row_idx, column=4, value=biz["website"])
        website_cell.border = thin_border
        website_cell.alignment = Alignment(horizontal="center")
        if biz["website"] == "Yes":
            website_cell.fill = yes_fill
        else:
            website_cell.fill = no_fill

    # Column widths
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 12

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:D{len(results) + 1}"

    # Save
    safe_type = re.sub(r"[^\w\s-]", "", business_type).strip().replace(" ", "_")
    safe_city = re.sub(r"[^\w\s-]", "", city).strip().replace(" ", "_")
    filename = f"{safe_type}_{safe_city}_leads.xlsx"
    output_dir = Path(__file__).parent / "output"
    output_dir.mkdir(exist_ok=True)
    filepath = output_dir / filename
    wb.save(filepath)
    return str(filepath)


def main():
    if len(sys.argv) == 3:
        business_type = sys.argv[1]
        city = sys.argv[2]
    elif len(sys.argv) == 1:
        print("=== Lead Scraper ===")
        business_type = input("Business type (e.g. plumbers): ").strip()
        city = input("City (e.g. Hammond Louisiana): ").strip()
    else:
        print("Usage: python scraper.py \"plumbers\" \"Hammond Louisiana\"")
        sys.exit(1)

    if not business_type or not city:
        print("Error: Both business type and city are required.")
        sys.exit(1)

    results = scrape_google_maps(business_type, city)

    if not results:
        print("\nNo results found. Google may have changed their layout or blocked the request.")
        print("Try again in a few minutes.")
        sys.exit(0)

    filepath = save_to_excel(results, business_type, city)

    # Summary
    with_site = sum(1 for r in results if r["website"] == "Yes")
    without_site = sum(1 for r in results if r["website"] == "No")

    print(f"\n{'='*50}")
    print(f"DONE! Found {len(results)} businesses")
    print(f"  With website:    {with_site}")
    print(f"  Without website: {without_site}  <-- your leads!")
    print(f"\nSaved to: {filepath}")
    print(f"{'='*50}")


if __name__ == "__main__":
    main()
